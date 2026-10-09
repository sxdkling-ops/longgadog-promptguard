"""Exercise the file reader with one generated sample per format. Fictional data only."""
import base64
import email.message
import gzip
import io
import os
import sys
import tarfile
import zipfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "gateway"))
import docx  # noqa: E402
import openpyxl  # noqa: E402
import pymupdf  # noqa: E402
from PIL import Image  # noqa: E402

import files  # noqa: E402
import store  # noqa: E402

SECRET_NAME = "Maria Santos"
BODY = "Patient: Maria Santos\nTIN 123-456-789-000\nDiagnosis: Acute bronchitis\nserver 10.20.30.44"
samples = {}


def page_png(lines, size=15):
    d = pymupdf.open()
    p = d.new_page(width=700, height=420)
    for i, ln in enumerate(lines):
        p.insert_text((40, 60 + i * 34), ln, fontsize=size)
    return p.get_pixmap(dpi=150).tobytes("png")


png = page_png(["Customer record", "Customer: Maria Santos", "Email: maria.santos@example.com", "Mobile: 0917 555 0142"])
pil = Image.open(io.BytesIO(png)).convert("RGB")


def as_format(fmt, **kw):
    b = io.BytesIO()
    pil.save(b, fmt, **kw)
    return b.getvalue()


samples["plain.txt"] = BODY.encode()
samples["script.py"] = b"# owner: Maria Santos\nAPI_KEY = 'pgdemo_k9Qz7Lm2Xv8Tn4Rb6Wd1Hs5Y'\nprint('ok')\n"
samples["data.json"] = b'{"customer": "x", "email": "maria.santos@example.com", "blob": "' + base64.b64encode(os.urandom(300)) + b'"}'
samples["legacy-win.txt"] = "Se\xf1or kay Maria Santos, TIN 123-456-789-000".encode("cp1252")
samples["renamed.dat"] = BODY.encode()                       # text with a misleading extension
samples["people.csv"] = b"Borrower,Birthdate,Salary,Remarks\nMaria Santos,1999-03-04,45000,call 0917 555 0142\n"

d = docx.Document()
d.add_paragraph("Atty. Ramon Villanueva, email ramon@example.com")
d.sections[0].header.paragraphs[0].text = "Confidential: client Northwind Lending Corp"
d.sections[0].footer.paragraphs[0].text = "Prepared for Mr. Benjie Soriano"
b = io.BytesIO(); d.save(b); samples["letter.docx"] = b.getvalue()
samples["letter-renamed.bin"] = b.getvalue()                 # docx with a wrong extension

wb = openpyxl.Workbook(); ws = wb.active
ws.append(["Borrower", "TIN", "Branch"]); ws.append(["Maria Santos", "123-456-789-000", "Pasig"])
ws2 = wb.create_sheet("Payroll - Santos"); ws2.append(["Employee", "Salary"]); ws2.append(["Paolo Mendoza", 38500])
b = io.BytesIO(); wb.save(b); samples["book.xlsx"] = b.getvalue()

b = io.BytesIO()
with zipfile.ZipFile(b, "w") as z:
    z.writestr("[Content_Types].xml", "<Types/>")
    z.writestr("ppt/presentation.xml", "<p/>")
    z.writestr("ppt/slides/slide1.xml", '<p:sld><a:p><a:r><a:t>Client: Maria Santos</a:t></a:r></a:p><a:p><a:r><a:t>TIN 123-456-789-000</a:t></a:r></a:p></p:sld>')
    z.writestr("ppt/media/image1.png", png)
samples["deck.pptx"] = b.getvalue()

b = io.BytesIO()
with zipfile.ZipFile(b, "w") as z:
    z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
    z.writestr("content.xml", "<office:document-content><office:body><office:text><text:p>Patient: Maria Santos</text:p><text:p>TIN 123-456-789-000</text:p></office:text></office:body></office:document-content>")
samples["notes.odt"] = b.getvalue()

samples["memo.rtf"] = rb"{\rtf1\ansi{\fonttbl{\f0 Arial;}}\f0\fs24 Patient: Maria Santos\par TIN 123-456-789-000\par Se\'f1ora}"

doc = pymupdf.open(); pg = doc.new_page(); pg.insert_text((72, 72), BODY)
samples["text.pdf"] = doc.tobytes()
doc = pymupdf.open(); pg = doc.new_page(width=700, height=420); pg.insert_image(pg.rect, stream=png)
samples["scanned.pdf"] = doc.tobytes()

samples["shot.png"] = png
samples["shot.jpg"] = as_format("JPEG", quality=90)
samples["shot.webp"] = as_format("WEBP")
samples["shot.bmp"] = as_format("BMP")
samples["shot.gif"] = as_format("GIF")
samples["shot.tiff"] = as_format("TIFF")
try:
    samples["shot.heic"] = as_format("HEIF")
except Exception as e:
    print("(could not generate HEIC sample:", type(e).__name__, ")")
samples["photo-renamed.txt"] = png                             # image with a wrong extension
samples["blank.png"] = page_png([])
samples["PASSPORT.jpg"] = page_png(["REPUBLIKA NG PILIPINAS", "PASAPORTE / PASSPORT", "Surname: DELA CRUZ"])

msg = email.message.EmailMessage()
msg["From"] = "Liza Reyes <liza.reyes@example.com>"; msg["To"] = "hr@example.org"; msg["Subject"] = "Loan of Maria Santos"
msg.set_content("Please see attached. TIN 123-456-789-000")
msg.add_attachment(samples["people.csv"], maintype="text", subtype="csv", filename="people.csv")
msg.add_attachment(samples["letter.docx"], maintype="application", subtype="octet-stream", filename="letter.docx")
samples["mail.eml"] = msg.as_bytes()

b = io.BytesIO()
with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("a/plain.txt", samples["plain.txt"]); z.writestr("a/book.xlsx", samples["book.xlsx"]); z.writestr("shot.jpg", samples["shot.jpg"])
samples["bundle.zip"] = b.getvalue()
b = io.BytesIO()
with zipfile.ZipFile(b, "w") as z:
    z.writestr("ok.txt", b"hello"); z.writestr("PASSPORT.jpg", samples["PASSPORT.jpg"])
samples["bundle-with-id.zip"] = b.getvalue()
samples["notes.txt.gz"] = gzip.compress(BODY.encode())
b = io.BytesIO()
with tarfile.open(fileobj=b, mode="w") as t:
    info = tarfile.TarInfo("plain.txt"); info.size = len(BODY); t.addfile(info, io.BytesIO(BODY.encode()))
samples["backup.tar"] = b.getvalue()

samples["random.bin"] = os.urandom(5000)
samples["video.mp4"] = b"\x00\x00\x00\x18ftypmp42" + os.urandom(4000)
samples["app.exe"] = b"MZ\x90\x00" + b"\x00" * 300 + b"This program cannot be run in DOS mode.\x00\x00Owner: Maria Santos\x00" + b"\x01\x02\x03" * 200
doc = pymupdf.open(); doc.new_page().insert_text((72, 72), "secret"); samples["locked.pdf"] = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")

sess = store.session("formats-test-1")
leaks = 0
for name, data in samples.items():
    r = files.process_file(name, data, sess, ["Northwind Lending Corp"], False)
    if r["blocked"]:
        print(f"BLOCKED  {name:<22} {r['reason']}")
        continue
    out = base64.b64decode(r["out_b64"]).decode("utf-8", "replace")
    leak = [w for w in ("Maria", "Santos", "123-456-789-000", "maria.santos@", "0917", "Villanueva", "Soriano", "Northwind", "Mendoza", "10.20.30.44", "sk_live", "Liza", "Payroll") if w.lower() in out.lower()]
    leaks += bool(leak)
    print(f"{'LEAK    ' if leak else 'OK      '} {name:<22} [{r['kind']}] -> {r['out_name']} {r['counts']}" + (f"  LEAKED: {leak}" if leak else ""))
    if leak or "-v" in sys.argv:
        print("         " + out[:300].replace("\n", " / "))
print("files with leaks:", leaks)
